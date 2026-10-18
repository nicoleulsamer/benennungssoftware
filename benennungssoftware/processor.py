from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
import re
import shutil

from .config import AppConfig, Project
from .text_extraction import extract_document_text


@dataclass(frozen=True)
class ProcessResult:
    source: Path
    target: Path
    status: str
    project_code: str | None = None
    reason: str | None = None


def process_scan_folder(
    config: AppConfig,
    *,
    dry_run: bool = False,
    today: date | None = None,
    log_path: Path | None = None,
) -> list[ProcessResult]:
    current_date = today or date.today()
    config.scan_folder.mkdir(parents=True, exist_ok=True)
    config.projects_root.mkdir(parents=True, exist_ok=True)
    config.unassigned_folder.mkdir(parents=True, exist_ok=True)

    results: list[ProcessResult] = []
    for source in sorted(config.scan_folder.iterdir()):
        if not source.is_file() or source.suffix.lower() not in config.allowed_extensions:
            continue

        project = match_project(source, config)
        if project is None:
            target = unique_path(config.unassigned_folder / source.name)
            result = _move(source, target, dry_run=dry_run, status="unassigned", reason="no_project_match")
            if not dry_run:
                _append_to_excel_workbook(config, result)
            results.append(result)
            continue

        target_folder = config.projects_root / project.folder
        filename = build_filename(config, project, source, current_date)
        target = unique_path(target_folder / filename)
        result = _move(source, target, dry_run=dry_run, status="assigned", project_code=project.code)
        if not dry_run:
            _append_to_excel_workbook(config, result)
        results.append(result)

    _write_rename_log(log_path or config.rename_log_path, results, dry_run=dry_run)
    return results


def process_single_file(
    config: AppConfig,
    source: Path,
    *,
    dry_run: bool = False,
    today: date | None = None,
    log_path: Path | None = None,
) -> list[ProcessResult]:
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in config.allowed_extensions:
        raise ValueError(f"Unsupported file extension: {source.suffix}")

    current_date = today or date.today()
    config.projects_root.mkdir(parents=True, exist_ok=True)
    config.unassigned_folder.mkdir(parents=True, exist_ok=True)

    project = match_project(source, config)
    if project is None:
        target = unique_path(config.unassigned_folder / source.name)
        result = _move(source, target, dry_run=dry_run, status="unassigned", reason="no_project_match")
        if not dry_run:
            _append_to_excel_workbook(config, result)
        _write_rename_log(log_path or config.rename_log_path, [result], dry_run=dry_run)
        return [result]

    target_folder = config.projects_root / project.folder
    filename = build_filename(config, project, source, current_date)
    target = unique_path(target_folder / filename)
    result = _move(source, target, dry_run=dry_run, status="assigned", project_code=project.code)
    if not dry_run:
        _append_to_excel_workbook(config, result)
    _write_rename_log(log_path or config.rename_log_path, [result], dry_run=dry_run)
    return [result]


def assign_unassigned_document(
    config: AppConfig,
    source: Path,
    project_code: str,
    *,
    dry_run: bool = False,
    today: date | None = None,
    log_path: Path | None = None,
) -> ProcessResult:
    project = find_project_by_code(config, project_code)
    if project is None:
        raise ValueError(f"Unknown project code: {project_code}")
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in config.allowed_extensions:
        raise ValueError(f"Unsupported file extension: {source.suffix}")

    current_date = today or date.today()
    target_folder = config.projects_root / project.folder
    filename = build_filename(config, project, source, current_date)
    target = unique_path(target_folder / filename)
    result = _move(source, target, dry_run=dry_run, status="manually_assigned", project_code=project.code)
    if not dry_run:
        _append_to_excel_workbook(config, result)
    _write_rename_log(log_path or config.rename_log_path, [result], dry_run=dry_run)
    return result


def list_unassigned_documents(config: AppConfig) -> list[Path]:
    if not config.unassigned_folder.exists():
        return []
    return [
        source
        for source in sorted(config.unassigned_folder.iterdir())
        if source.is_file() and source.suffix.lower() in config.allowed_extensions
    ]


def find_project_by_code(config: AppConfig, project_code: str) -> Project | None:
    normalized = project_code.casefold()
    for project in config.projects:
        if project.code.casefold() == normalized:
            return project
    return None


def suggest_projects(source: Path, config: AppConfig, limit: int = 5) -> list[tuple[Project, int]]:
    text = extract_document_text(source, config.text_extraction)
    haystack = f"{source.stem} {text}".casefold()
    matches = _matching_projects(haystack, config.projects)
    if not matches:
        return []

    ranked = sorted(
        ((project, _project_match_score(haystack, project)) for project in matches),
        key=lambda entry: (-entry[1], entry[0].code.casefold()),
    )
    return [(project, score) for project, score in ranked[:limit]]


def match_project(source: Path, config: AppConfig) -> Project | None:
    suggestions = suggest_projects(source, config, limit=10)
    if not suggestions:
        filename_matches = _matching_projects(source.stem.casefold(), config.projects)
        if len(filename_matches) == 1:
            return filename_matches[0]
        return None

    best_project, best_score = suggestions[0]
    if len(suggestions) > 1:
        second_score = suggestions[1][1]
        if second_score == best_score:
            return None

    if best_score <= 0:
        return None

    return best_project


def _matching_projects(haystack: str, projects: tuple[Project, ...]) -> list[Project]:
    normalized_haystack = re.sub(r"[^a-z0-9]+", " ", haystack.casefold()).strip()
    return [
        project
        for project in projects
        if any(
            keyword and re.sub(r"[^a-z0-9]+", " ", keyword.casefold()).strip() in normalized_haystack
            for keyword in project.keywords
        )
    ]


GENERIC_PROJECT_KEYWORDS = {
    "wien",
    "innsbruck",
    "hall",
    "vols",
    "lans",
    "st",
    "au",
    "gasse",
    "str",
    "strae",
    "strasse",
    "weg",
    "platz",
    "platzl",
    "johann",
    "schon",
    "hotting",
    "muhlau",
    "birgitz",
    "seefeld",
    "kufstein",
    "rum",
    "arzl",
    "westendorf",
    "mutters",
    "jenbach",
    "schwaz",
    "fieberbrunn",
    "pradl",
    "hall",
    "sistrans",
    "puitnegg",
    "kranebitten",
    "alpenhof",
    "dokument",
    "rechnung",
    "scan",
}


def _keyword_weight(normalized_keyword: str) -> int:
    if not normalized_keyword:
        return 0
    if normalized_keyword in GENERIC_PROJECT_KEYWORDS:
        return 1
    if " " in normalized_keyword:
        return 12
    if len(normalized_keyword) <= 3:
        return 1
    if len(normalized_keyword) <= 5:
        return 3
    return 5


def _project_match_score(haystack: str, project: Project) -> int:
    normalized_haystack = re.sub(r"[^a-z0-9]+", " ", haystack.casefold()).strip()
    score = 0
    seen_hits: set[str] = set()

    for keyword in project.keywords:
        normalized_keyword = re.sub(r"[^a-z0-9]+", " ", keyword.casefold()).strip()
        if not normalized_keyword:
            continue
        if normalized_keyword in normalized_haystack:
            weight = _keyword_weight(normalized_keyword)
            score += weight
            seen_hits.add(normalized_keyword)

    if project.code.casefold() in haystack.casefold():
        score += 15

    if len(seen_hits) > 1:
        score += min(len(seen_hits) * 3, 18)

    if len(seen_hits) == 1 and next(iter(seen_hits)) in GENERIC_PROJECT_KEYWORDS:
        score = max(score - 3, 0)

    return score


def parse_company_entries(text: str) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    current: dict[str, str] | None = None

    def consume_field(line: str) -> tuple[str, str] | None:
        cleaned = re.sub(r"^(?:[A-Za-zÄÖÜäöüß]{1,3}\.?\s+)+", "", line.strip())
        field_match = re.match(
            r"^(UID-Nummer|UID-Nr|UID|Firmenbuchnummer|Firmenbuch|Gründung|Steuernummer|Steuernr)\s*[:.-]?\s*(.*)$",
            cleaned,
            flags=re.IGNORECASE,
        )
        if field_match is None:
            return None

        label, value = field_match.groups()
        key_map = {
            "uid-nummer": "uid_nummer",
            "uid-nr": "uid_nummer",
            "uid": "uid_nummer",
            "firmenbuchnummer": "firmenbuchnummer",
            "firmenbuch": "firmenbuchnummer",
            "gründung": "gruendung",
            "gruendung": "gruendung",
            "steuernummer": "steuernummer",
            "steuernr": "steuernummer",
        }
        normalized_value = value.strip() or "folgt"
        normalized_label = label.casefold()
        if normalized_label not in key_map:
            return None
        return key_map[normalized_label], normalized_value

    for raw_line in text.splitlines():
        line = " ".join(raw_line.split())
        if not line:
            continue
        if line.startswith("FIRMENLISTE"):
            if current is not None:
                entries.append(current)
            break

        entry_match = re.match(r"^\((?P<tag>[A-Z0-9]+)\)\s*(?P<body>.+)$", line)
        if entry_match:
            if current is not None:
                entries.append(current)
            body = entry_match.group("body").strip()
            address_match = re.search(r"(?P<address>\d{4}\s+[A-Za-zÄÖÜäöüß0-9.,/ -]+)$", body)
            if address_match:
                company_name = body[: address_match.start()].strip().rstrip(" ,;-")
                address = address_match.group("address").strip()
            else:
                company_name = body.strip()
                address = ""
            current = {
                "kennzeichen": entry_match.group("tag"),
                "firma": company_name,
                "adresse": address,
            }
            continue

        if current is None:
            continue

        field = consume_field(line)
        if field is not None:
            current[field[0]] = field[1]

    if current is not None:
        entries.append(current)

    return entries


def build_filename(config: AppConfig, project: Project, source: Path, current_date: date) -> str:
    document_label = _document_label(source)
    if document_label is not None:
        return f"{document_label}{source.suffix.lower()}"

    values = {
        "date": current_date.isoformat(),
        "project_code": sanitize(project.code),
        "document_type": sanitize(config.default_document_type),
        "original_stem": sanitize(source.stem),
        "document_label": sanitize(source.stem),
        "extension": source.suffix.lower(),
    }
    return config.name_schema.format(**values)


def _document_label(source: Path) -> str | None:
    stem = source.stem.strip()
    if not stem:
        return None

    if re.match(r"^[A-Z]{2,}(?:[_-][A-Za-z0-9]+)+$", stem):
        return sanitize(stem)
    return None


def sanitize(value: str) -> str:
    transliterated = (
        value.strip()
        .replace("ä", "ae")
        .replace("ö", "oe")
        .replace("ü", "ue")
        .replace("Ä", "Ae")
        .replace("Ö", "Oe")
        .replace("Ü", "Ue")
        .replace("ß", "ss")
    )
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", transliterated)
    cleaned = re.sub(r"-{2,}", "-", cleaned).strip("-._")
    return cleaned or "unbenannt"


def unique_path(target: Path) -> Path:
    if not target.exists():
        return target

    stem = target.stem
    suffix = target.suffix
    parent = target.parent
    counter = 1
    while True:
        candidate = parent / f"{stem}_{counter:03d}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def _move(
    source: Path,
    target: Path,
    *,
    dry_run: bool,
    status: str,
    project_code: str | None = None,
    reason: str | None = None,
) -> ProcessResult:
    if not dry_run:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source), str(target))
    return ProcessResult(source=source, target=target, status=status, project_code=project_code, reason=reason)


def _write_rename_log(path: Path | None, results: list[ProcessResult], *, dry_run: bool) -> None:
    if path is None:
        return
    from .reporting import write_process_report

    write_process_report(path, results, dry_run=dry_run)


def _append_to_excel_workbook(config: AppConfig, result: ProcessResult) -> None:
    if not config.excel.enabled or config.excel.workbook_path is None:
        return

    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return

    workbook_path = config.excel.workbook_path
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()

    if config.excel.sheet_name not in workbook.sheetnames:
        worksheet = workbook.active
        worksheet.title = config.excel.sheet_name
    else:
        worksheet = workbook[config.excel.sheet_name]

    if worksheet.max_row == 1 and worksheet["A1"].value is None:
        worksheet.delete_rows(1, 1)
        worksheet.append(list(config.excel.columns))
    elif worksheet.max_row == 0:
        worksheet.append(list(config.excel.columns))
    elif worksheet["A1"].value != config.excel.columns[0]:
        worksheet.insert_rows(1)
        worksheet["A1"] = config.excel.columns[0]
        for column_index, header in enumerate(config.excel.columns[1:], start=2):
            worksheet.cell(row=1, column=column_index, value=header)

    next_number = worksheet.max_row
    if worksheet.max_row == 1 and worksheet["A1"].value == config.excel.columns[0]:
        next_number = 1

    worksheet.append([
        next_number,
        result.source.name,
        result.target.name,
        result.status,
        result.reason or "",
    ])
    workbook.save(workbook_path)


def append_company_entries_to_excel(config: AppConfig, entries: list[dict[str, str]]) -> None:
    if not config.excel.enabled or config.excel.workbook_path is None or not entries:
        return

    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return

    workbook_path = config.excel.workbook_path
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()

    if config.excel.sheet_name not in workbook.sheetnames:
        worksheet = workbook.active
        worksheet.title = config.excel.sheet_name
    else:
        worksheet = workbook[config.excel.sheet_name]

    headers = list(config.excel.columns)
    if worksheet.max_row == 0:
        worksheet.append(headers)
    elif worksheet.max_row == 1 and worksheet["A1"].value is None:
        worksheet.delete_rows(1, 1)
        worksheet.append(headers)
    elif worksheet["A1"].value != headers[0]:
        worksheet.insert_rows(1)
        worksheet["A1"] = headers[0]
        for column_index, header in enumerate(headers[1:], start=2):
            worksheet.cell(row=1, column=column_index, value=header)

    next_number = worksheet.max_row
    if worksheet.max_row == 1 and worksheet["A1"].value == headers[0]:
        next_number = 1

    for entry in entries:
        row = [next_number]
        for column in headers[1:]:
            value = ""
            normalized_column = column.casefold().replace(" ", "-")
            if normalized_column in {"firma", "firma-name", "unternehmen"}:
                value = entry.get("firma", "")
            elif normalized_column in {"kennzeichen", "code", "tag"}:
                value = entry.get("kennzeichen", "")
            elif normalized_column in {"adresse", "ort", "anschrift"}:
                value = entry.get("adresse", "")
            elif normalized_column in {"uid-nummer", "uid"}:
                value = entry.get("uid_nummer", "")
            elif normalized_column in {"firmenbuchnummer", "firm-number", "firmenbuch"}:
                value = entry.get("firmenbuchnummer", "")
            elif normalized_column in {"gruendung", "gründung", "gegruendet"}:
                value = entry.get("gruendung", "")
            elif normalized_column in {"status", "zustand", "erfasst"}:
                value = "erfasst"
            elif normalized_column in {"hinweis", "notiz"}:
                value = entry.get("hinweis", "")
            row.append(value)
        worksheet.append(row)
        next_number += 1

    workbook.save(workbook_path)


def import_company_list_pdf(config: AppConfig, source: Path) -> list[dict[str, str]]:
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(source)

    text = extract_document_text(source, config.text_extraction)
    entries = parse_company_entries(text)
    if not entries:
        return []

    append_company_entries_to_excel(config, entries)
    return entries


def parse_bkp_entries(text: str) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    pattern = re.compile(r"(?<!\d)(\d{4})\s+([^\d][^\n]*?)(?=\s+\d{4}\s+|$)")

    for raw_line in text.splitlines():
        line = " ".join(raw_line.split())
        if not line:
            continue
        if "BKP" in line.upper() or line.upper().startswith("BAUKOSTEN"):
            continue

        matches = pattern.findall(line)
        for code, description in matches:
            cleaned = description.strip()
            if not cleaned:
                continue
            entries.append({
                "code": code.strip(),
                "beschreibung": cleaned,
            })
    return entries


def append_bkp_entries_to_excel(config: AppConfig, entries: list[dict[str, str]]) -> None:
    if not config.excel.enabled or config.excel.workbook_path is None or not entries:
        return

    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        return

    workbook_path = config.excel.workbook_path
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
    else:
        workbook = Workbook()

    if config.excel.sheet_name not in workbook.sheetnames:
        worksheet = workbook.active
        worksheet.title = config.excel.sheet_name
    else:
        worksheet = workbook[config.excel.sheet_name]

    headers = list(config.excel.columns)
    if worksheet.max_row == 0:
        worksheet.append(headers)
    elif worksheet.max_row == 1 and worksheet["A1"].value is None:
        worksheet.delete_rows(1, 1)
        worksheet.append(headers)
    elif worksheet["A1"].value != headers[0]:
        worksheet.insert_rows(1)
        worksheet["A1"] = headers[0]
        for column_index, header in enumerate(headers[1:], start=2):
            worksheet.cell(row=1, column=column_index, value=header)

    next_number = worksheet.max_row
    if worksheet.max_row == 1 and worksheet["A1"].value == headers[0]:
        next_number = 1

    for entry in entries:
        row = [next_number]
        for column in headers[1:]:
            value = ""
            normalized_column = column.casefold().replace(" ", "-")
            if normalized_column in {"code", "nummer", "kontennummer"}:
                value = entry.get("code", "")
            elif normalized_column in {"beschreibung", "bezeichnung", "name"}:
                value = entry.get("beschreibung", "")
            row.append(value)
        worksheet.append(row)
        next_number += 1

    workbook.save(workbook_path)


def import_bkp_list_pdf(config: AppConfig, source: Path) -> list[dict[str, str]]:
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(source)

    text = extract_document_text(source, config.text_extraction)
    entries = parse_bkp_entries(text)
    if not entries:
        return []

    append_bkp_entries_to_excel(config, entries)
    return entries


def _normalize_sheet_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")


def import_company_list_workbook(config: AppConfig, workbook_path: Path) -> list[dict[str, str]]:
    if not workbook_path.exists() or not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to import Excel files") from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    target_sheet = None
    candidate_names = [config.excel.sheet_name, "Projektliste", "Firmenliste", "Projektübersicht", "Firmenübersicht"]
    for name in candidate_names:
        if name in workbook.sheetnames:
            target_sheet = workbook[name]
            break
    if target_sheet is None:
        target_sheet = workbook[workbook.sheetnames[0]]

    rows = list(target_sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    header_map: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        if cell is None:
            continue
        header_map[_normalize_sheet_header(str(cell))] = index

    field_aliases = {
        "firma": "firma",
        "unternehmen": "firma",
        "projekt": "firma",
        "name": "firma",
        "kennzeichen": "kennzeichen",
        "code": "kennzeichen",
        "adresse": "adresse",
        "anschrift": "adresse",
        "ort": "adresse",
        "uid": "uid_nummer",
        "uid_nummer": "uid_nummer",
        "uidnummer": "uid_nummer",
        "firmenbuchnummer": "firmenbuchnummer",
        "firmenbuch": "firmenbuchnummer",
        "gruendung": "gruendung",
        "gründung": "gruendung",
        "gegrundet": "gruendung",
        "status": "status",
    }

    column_indexes = {}
    for canonical, alias in field_aliases.items():
        index = header_map.get(canonical)
        if index is None:
            continue
        column_indexes[alias] = index

    entries: list[dict[str, str]] = []
    for row in rows[1:]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        entry: dict[str, str] = {}
        if 0 in range(len(row)):
            pass
        for field_name, index in column_indexes.items():
            if index >= len(row):
                continue
            value = row[index]
            if value is None:
                value = ""
            entry[field_name] = str(value).strip()
        if not entry:
            continue
        if "firma" not in entry and row and row[0] is not None:
            entry["firma"] = str(row[0]).strip()
        if "status" not in entry:
            entry["status"] = "erfasst"
        entries.append(entry)

    try:
        if entries:
            append_company_entries_to_excel(config, entries)
    finally:
        workbook.close()
    return entries
