from __future__ import annotations

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from benennungssoftware.config import AppConfig, ExcelConfig, Project, load_config
from benennungssoftware.processor import (
    append_company_entries_to_excel,
    assign_unassigned_document,
    import_company_list_pdf,
    import_company_list_workbook,
    list_unassigned_documents,
    match_project,
    parse_company_entries,
    process_scan_folder,
    process_single_file,
    sanitize,
)


class ProcessorTests(unittest.TestCase):
    def test_assigns_matching_scan_to_project_folder(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            projects_root = root / "projects"
            unassigned_folder = root / "unassigned"
            scan_folder.mkdir()
            source = scan_folder / "rechnung-kunde-a.pdf"
            source.write_bytes(b"pdf")

            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=projects_root,
                unassigned_folder=unassigned_folder,
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Rechnung",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde-a", "kunde a")),),
            )

            results = process_scan_folder(config, today=date(2026, 6, 14))

            self.assertEqual(len(results), 1)
            self.assertEqual(results[0].status, "assigned")
            self.assertFalse(source.exists())
            self.assertTrue((projects_root / "PRJ001_Test" / "2026-06-14_PRJ001_Rechnung_rechnung-kunde-a.pdf").exists())

    def test_writes_default_rename_log_for_processed_file(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            scan_folder.mkdir()
            source = scan_folder / "rechnung-kunde-a.pdf"
            source.write_bytes(b"pdf")
            log_path = root / "protokoll_umbenennung.csv"

            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Rechnung",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde-a", "kunde a")),),
                rename_log_path=log_path,
            )

            process_scan_folder(config, today=date(2026, 6, 14))

            self.assertTrue(log_path.exists())
            content = log_path.read_text(encoding="utf-8")
            self.assertIn("assigned", content)
            self.assertIn("rechnung-kunde-a.pdf", content)

    def test_unassigned_scan_is_moved_to_unassigned_folder(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            scan_folder.mkdir()
            source = scan_folder / "scan-ohne-treffer.txt"
            source.write_text("Kein Projektbezug", encoding="utf-8")

            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".txt",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("anderes",)),),
            )

            results = process_scan_folder(config)

            self.assertEqual(results[0].status, "unassigned")
            self.assertEqual(results[0].reason, "no_project_match")
            self.assertTrue((root / "unassigned" / "scan-ohne-treffer.txt").exists())

    def test_existing_target_is_not_overwritten(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            target_folder = root / "projects" / "PRJ001_Test"
            scan_folder.mkdir()
            target_folder.mkdir(parents=True)
            source = scan_folder / "kunde-a.pdf"
            source.write_bytes(b"new")
            existing = target_folder / "2026-06-14_PRJ001_Dokument_kunde-a.pdf"
            existing.write_bytes(b"existing")

            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde-a",)),),
            )

            process_scan_folder(config, today=date(2026, 6, 14))

            self.assertEqual(existing.read_bytes(), b"existing")
            self.assertTrue((target_folder / "2026-06-14_PRJ001_Dokument_kunde-a_001.pdf").exists())

    def test_assigns_pdf_by_extracted_text(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            scan_folder.mkdir()
            source = scan_folder / "scan001.pdf"
            source.write_bytes(b"pdf")

            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde a",)),),
            )

            with patch("benennungssoftware.processor.extract_document_text", return_value="Rechnung fuer Kunde A"):
                results = process_scan_folder(config, today=date(2026, 6, 14))

            self.assertEqual(results[0].status, "assigned")
            self.assertTrue((root / "projects" / "PRJ001_Test" / "2026-06-14_PRJ001_Dokument_scan001.pdf").exists())

    def test_real_demo_invoice_matches_project_list(self) -> None:
        root = Path(__file__).resolve().parents[1]
        config = load_config(root / "examples" / "config.json")
        source = root / "examples" / "demo" / "scans" / "Scan_08_04_26_16_26_52.pdf"

        project = match_project(source, config)

        self.assertIsNotNone(project)
        self.assertEqual(project.code, "P-54")

    def test_manually_assigns_unassigned_document(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            unassigned_folder = root / "unassigned"
            unassigned_folder.mkdir()
            source = unassigned_folder / "scan003_unbekannt.txt"
            source.write_text("Kein Treffer", encoding="utf-8")

            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=unassigned_folder,
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".txt",),
                projects=(Project(code="PRJ002", folder="PRJ002_Test", keywords=("kunde b",)),),
            )

            result = assign_unassigned_document(config, source, "PRJ002", today=date(2026, 6, 14))

            self.assertEqual(result.status, "manually_assigned")
            self.assertFalse(source.exists())
            self.assertTrue((root / "projects" / "PRJ002_Test" / "2026-06-14_PRJ002_Dokument_scan003_unbekannt.txt").exists())

    def test_manual_assignment_rejects_unknown_project_code(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "scan.txt"
            source.write_text("Text", encoding="utf-8")
            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".txt",),
                projects=(),
            )

            with self.assertRaises(ValueError):
                assign_unassigned_document(config, source, "PRJ999")

    def test_lists_unassigned_documents_with_allowed_extensions(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            unassigned_folder = root / "unassigned"
            unassigned_folder.mkdir()
            first = unassigned_folder / "a.txt"
            second = unassigned_folder / "b.pdf"
            ignored = unassigned_folder / "ignore.tmp"
            first.write_text("A", encoding="utf-8")
            second.write_bytes(b"pdf")
            ignored.write_text("tmp", encoding="utf-8")

            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=unassigned_folder,
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".pdf", ".txt"),
                projects=(),
            )

            self.assertEqual(list_unassigned_documents(config), [first, second])

    def test_lists_no_unassigned_documents_when_folder_is_missing(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".pdf", ".txt"),
                projects=(),
            )

            self.assertEqual(list_unassigned_documents(config), [])

    def test_processes_single_file_for_gui_workflow(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "rechnung-kunde-a.pdf"
            source.write_bytes(b"pdf")
            config = AppConfig(
                scan_folder=root / "scan",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Dokument",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde a",)),),
            )

            results = process_single_file(config, source, today=date(2026, 6, 14))

            self.assertEqual(results[0].status, "assigned")
            self.assertFalse(source.exists())
            self.assertTrue((root / "projects" / "PRJ001_Test" / "2026-06-14_PRJ001_Dokument_rechnung-kunde-a.pdf").exists())

    def test_processes_single_file_and_updates_excel_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            scan_folder = root / "scans"
            scan_folder.mkdir()
            source = scan_folder / "rechnung-kunde-a.pdf"
            source.write_bytes(b"pdf")
            workbook_path = root / "Rechnungsuebersicht.xlsx"
            config = AppConfig(
                scan_folder=scan_folder,
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Rechnung",
                allowed_extensions=(".pdf",),
                projects=(Project(code="PRJ001", folder="PRJ001_Test", keywords=("kunde a",)),),
                excel=ExcelConfig(
                    enabled=True,
                    workbook_path=workbook_path,
                    workbook_code="9543",
                    sheet_name="Rechnungsübersicht",
                    columns=("Nr", "Originaldatei", "Neuer Name", "Status", "Hinweis"),
                ),
            )

            results = process_single_file(config, source, today=date(2026, 6, 14))

            self.assertEqual(results[0].status, "assigned")
            self.assertTrue(workbook_path.exists())

            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path)
            sheet = workbook["Rechnungsübersicht"]
            self.assertEqual(sheet["A2"].value, 1)
            self.assertEqual(sheet["B2"].value, "rechnung-kunde-a.pdf")
            self.assertEqual(sheet["C2"].value, "2026-06-14_PRJ001_Rechnung_rechnung-kunde-a.pdf")
            self.assertEqual(sheet["D2"].value, "assigned")

    def test_parses_company_entries_from_pdf_text(self) -> None:
        text = """(MH) Moser Immobilien Holding GmbH 6020 Innsbruck, Sparkassenplatz 5
han UID-Nummer ATU74156637
GF Firmenbuchnummer FN 507097g
Gründung 2019
(MIB) Moser Immobilien Beteiligungs GmbH 6020 Innsbruck, Sparkassenplatz 5
han UID-Nummer
GF Firmenbuchnummer FN 652501k
Gründung 2025
"""

        entries = parse_company_entries(text)

        self.assertEqual(len(entries), 2)
        self.assertEqual(entries[0]["firma"], "Moser Immobilien Holding GmbH")
        self.assertEqual(entries[0]["uid_nummer"], "ATU74156637")
        self.assertEqual(entries[0]["firmenbuchnummer"], "FN 507097g")
        self.assertEqual(entries[0]["gruendung"], "2019")

    def test_parses_real_world_company_list_format(self) -> None:
        text = """(DAV) MIW DAV59 GmbH & Co KG 1010 Wien, Schottengasse 4/26
han. UID-Nummer ATU 77788878
GF Firmenbuchnummer FN573805g
Gründung 21.01.2022
(QU) QUANTUM Engagement GmbH 6020 Innsbruck, Sparkassenplatz 5
han UID-Nummer folgt
GF Firmenbuchnummer FN 564990v
Gründung 13.09.2021
"""

        entries = parse_company_entries(text)

        self.assertEqual(len(entries), 2)
        self.assertEqual(entries[0]["firma"], "MIW DAV59 GmbH & Co KG")
        self.assertEqual(entries[0]["uid_nummer"], "ATU 77788878")
        self.assertEqual(entries[0]["firmenbuchnummer"], "FN573805g")
        self.assertEqual(entries[1]["uid_nummer"], "folgt")

    def test_appends_company_entries_to_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "Firmenliste.xlsx"
            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Firma",
                allowed_extensions=(".pdf",),
                projects=(),
                excel=ExcelConfig(
                    enabled=True,
                    workbook_path=workbook_path,
                    workbook_code="9543",
                    sheet_name="Firmenliste",
                    columns=("Nr", "Firma", "Kennzeichen", "Adresse", "UID-Nummer", "Firmenbuchnummer", "Gründung", "Status"),
                ),
            )

            append_company_entries_to_excel(
                config,
                [{
                    "kennzeichen": "MH",
                    "firma": "Moser Immobilien Holding GmbH",
                    "adresse": "6020 Innsbruck, Sparkassenplatz 5",
                    "uid_nummer": "ATU74156637",
                    "firmenbuchnummer": "FN 507097g",
                    "gruendung": "2019",
                }],
            )

            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path)
            sheet = workbook["Firmenliste"]
            self.assertEqual(sheet["A2"].value, 1)
            self.assertEqual(sheet["B2"].value, "Moser Immobilien Holding GmbH")
            self.assertEqual(sheet["C2"].value, "MH")
            self.assertEqual(sheet["D2"].value, "6020 Innsbruck, Sparkassenplatz 5")
            self.assertEqual(sheet["E2"].value, "ATU74156637")
            self.assertEqual(sheet["F2"].value, "FN 507097g")
            self.assertEqual(sheet["G2"].value, "2019")
            self.assertEqual(sheet["H2"].value, "erfasst")

    def test_imports_company_list_pdf_into_excel(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "firmenliste.pdf"
            source.write_bytes(b"pdf")
            workbook_path = root / "Firmenliste.xlsx"
            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Firma",
                allowed_extensions=(".pdf",),
                projects=(),
                excel=ExcelConfig(
                    enabled=True,
                    workbook_path=workbook_path,
                    workbook_code="9543",
                    sheet_name="Firmenliste",
                    columns=("Nr", "Firma", "Kennzeichen", "Adresse", "UID-Nummer", "Firmenbuchnummer", "Gründung", "Status"),
                ),
            )

            with patch("benennungssoftware.processor.extract_document_text", return_value="""(MH) Moser Immobilien Holding GmbH 6020 Innsbruck, Sparkassenplatz 5
han UID-Nummer ATU74156637
GF Firmenbuchnummer FN 507097g
Gründung 2019
"""):
                entries = import_company_list_pdf(config, source)

            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]["firma"], "Moser Immobilien Holding GmbH")

            from openpyxl import load_workbook

            workbook = load_workbook(workbook_path)
            sheet = workbook["Firmenliste"]
            self.assertEqual(sheet["B2"].value, "Moser Immobilien Holding GmbH")
            self.assertEqual(sheet["E2"].value, "ATU74156637")
            self.assertEqual(sheet["H2"].value, "erfasst")

    def test_imports_company_list_from_excel_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "Projektliste.xlsx"
            workbook = __import__("openpyxl").Workbook()
            sheet = workbook.active
            sheet.title = "Projektliste"
            sheet.append(["Firma", "Kennzeichen", "Adresse", "UID-Nummer", "Firmenbuchnummer", "Gründung", "Status"])
            sheet.append(["Moser Immobilien Holding GmbH", "MH", "6020 Innsbruck, Sparkassenplatz 5", "ATU74156637", "FN 507097g", "2019", "erfasst"])
            workbook.save(workbook_path)
            workbook.close()

            config = AppConfig(
                scan_folder=root / "scans",
                projects_root=root / "projects",
                unassigned_folder=root / "unassigned",
                name_schema="{date}_{project_code}_{document_type}_{original_stem}{extension}",
                default_document_type="Firma",
                allowed_extensions=(".pdf",),
                projects=(),
                excel=ExcelConfig(
                    enabled=True,
                    workbook_path=root / "Firmenliste.xlsx",
                    workbook_code="9543",
                    sheet_name="Firmenliste",
                    columns=("Nr", "Firma", "Kennzeichen", "Adresse", "UID-Nummer", "Firmenbuchnummer", "Gründung", "Status"),
                ),
            )

            entries = import_company_list_workbook(config, workbook_path)

            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]["firma"], "Moser Immobilien Holding GmbH")
            self.assertEqual(entries[0]["uid_nummer"], "ATU74156637")

            from openpyxl import load_workbook
            result_wb = load_workbook(config.excel.workbook_path)
            result_sheet = result_wb["Firmenliste"]
            self.assertEqual(result_sheet["B2"].value, "Moser Immobilien Holding GmbH")
            self.assertEqual(result_sheet["E2"].value, "ATU74156637")

    def test_sanitize_removes_unsafe_filename_characters(self) -> None:
        self.assertEqual(sanitize(" Rechnung: Kunde A / 2026 "), "Rechnung-Kunde-A-2026")

    def test_sanitize_transliterates_german_characters(self) -> None:
        self.assertEqual(
            sanitize("Präsentation zur Studie Sexuelle Belästigung"),
            "Praesentation-zur-Studie-Sexuelle-Belaestigung",
        )
        self.assertEqual(sanitize("Grüße aus Österreich"), "Gruesse-aus-Oesterreich")


if __name__ == "__main__":
    unittest.main()
